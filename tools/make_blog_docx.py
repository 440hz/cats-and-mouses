"""Сборка article.docx для Блога RuStore из BLOG-ARTICLE.md (dev-only).

Требования RuStore: шрифт Arial, название статьи и все подзаголовки — стилем
«Заголовок 1», изображения в тексте на своих местах, без ручного форматирования.

Запуск: python tools/make_blog_docx.py
"""

import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

SRC = "store/blog/BLOG-ARTICLE.md"
OUT = "store/blog/article.docx"
COVER = "store/blog/01_обложка.png"
FONT = "Arial"

_TITLE_MARK = "# «Кошки-мышки»"
_STOP_MARK = "## Примечания по оформлению"

_RE_INLINE = re.compile(r"(\*\*.+?\*\*|\[[^\]]+\]\([^)]+\))")


def set_style_font(style, name, size=None):
    style.font.name = name
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)
    if size is not None:
        style.font.size = Pt(size)


def add_hyperlink(paragraph, url, text):
    r_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT)
    rpr.append(rfonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    run.append(rpr)
    txt = OxmlElement("w:t")
    txt.set(qn("xml:space"), "preserve")
    txt.text = text
    run.append(txt)
    link.append(run)
    paragraph._p.append(link)


def add_inline(paragraph, text):
    for part in _RE_INLINE.split(text):
        if not part:
            continue
        bold = re.fullmatch(r"\*\*(.+?)\*\*", part)
        if bold:
            run = paragraph.add_run(bold.group(1))
            run.bold = True
            run.font.name = FONT
            continue
        link = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", part)
        if link:
            add_hyperlink(paragraph, link.group(2), link.group(1))
            continue
        run = paragraph.add_run(part)
        run.font.name = FONT


def read_body():
    with open(SRC, encoding="utf-8") as fh:
        lines = fh.read().splitlines()
    start = next(i for i, ln in enumerate(lines) if ln.startswith(_TITLE_MARK))
    end = next(i for i, ln in enumerate(lines) if ln.startswith(_STOP_MARK))
    return lines[start:end]


def main():
    doc = Document()
    set_style_font(doc.styles["Normal"], FONT, 12)
    set_style_font(doc.styles["Heading 1"], FONT, 20)
    set_style_font(doc.styles["List Bullet"], FONT, 12)

    for raw in read_body():
        line = raw.strip()
        if not line or line == "---":
            continue
        if line.startswith("#"):
            p = doc.add_paragraph(style="Heading 1")
            add_inline(p, line.lstrip("#").strip())
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, line[2:])
        elif line.startswith("[ОБЛОЖКА"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(COVER, width=Inches(6.0))
        else:
            p = doc.add_paragraph()
            add_inline(p, line)

    doc.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
